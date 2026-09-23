from pathlib import Path
from tempfile import NamedTemporaryFile

import fitz
import pytesseract
from docx import Document
from fastapi import UploadFile
from PIL import Image
from pypdf import PdfReader

from app.text_cleaning import clean_text, word_count


class ExtractionError(RuntimeError):
    pass


def extract_from_path(path: Path) -> dict:
    suffix = path.suffix.lower()
    warnings: list[str] = []

    if suffix == ".pdf":
        text, method, pdf_warnings = _extract_pdf(path)
        warnings.extend(pdf_warnings)
    elif suffix == ".docx":
        text = _extract_docx(path)
        method = "docx"
    elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}:
        text = _ocr_image(path)
        method = "image_ocr"
    else:
        raise ExtractionError(f"format file tidak didukung: {suffix}")

    cleaned = clean_text(text)
    if not cleaned:
        raise ExtractionError("teks tidak dapat diekstrak dari file")

    return {
        "source": str(path),
        "file_type": suffix.lstrip("."),
        "extraction_method": method,
        "text": cleaned,
        "character_count": len(cleaned),
        "word_count": word_count(cleaned),
        "warnings": warnings,
    }


async def extract_from_upload(upload: UploadFile) -> dict:
    suffix = Path(upload.filename or "").suffix.lower()
    if not suffix:
        raise ExtractionError("nama file tidak memiliki ekstensi")

    with NamedTemporaryFile(suffix=suffix, delete=True) as tmp:
        while chunk := await upload.read(1024 * 1024):
            tmp.write(chunk)
        tmp.flush()
        return extract_from_path(Path(tmp.name))


def _extract_pdf(path: Path) -> tuple[str, str, list[str]]:
    warnings: list[str] = []
    text = _extract_pdf_digital(path)

    if len(clean_text(text)) >= 120:
        return text, "pdf_text", warnings

    warnings.append("teks digital PDF minim, OCR fallback digunakan")
    ocr_text = _ocr_pdf(path)
    return ocr_text, "pdf_ocr", warnings


def _extract_pdf_digital(path: Path) -> str:
    parts: list[str] = []

    try:
        reader = PdfReader(str(path))
        for page in reader.pages:
            parts.append(page.extract_text() or "")
    except Exception:
        parts = []

    if clean_text("\n".join(parts)):
        return "\n".join(parts)

    try:
        with fitz.open(str(path)) as document:
            return "\n".join(page.get_text("text") for page in document)
    except Exception as exc:
        raise ExtractionError(f"Gagal membaca PDF digital: {exc}") from exc


def _safe_ocr(image: Image.Image) -> str:
    """Mencoba OCR dengan ind+eng, lalu fallback ke eng atau default jika paket bahasa tertentu belum terpasang."""
    try:
        return pytesseract.image_to_string(image, lang="ind+eng")
    except Exception:
        try:
            return pytesseract.image_to_string(image, lang="eng")
        except Exception:
            return pytesseract.image_to_string(image)


def _ocr_pdf(path: Path) -> str:
    parts: list[str] = []
    try:
        with fitz.open(str(path)) as document:
            for page in document:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                parts.append(_safe_ocr(image))
    except Exception as exc:
        raise ExtractionError(f"Gagal memproses OCR PDF: {exc}") from exc
    return "\n".join(parts)


def _extract_docx(path: Path) -> str:
    try:
        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables: list[str] = []

        for table in document.tables:
            for row in table.rows:
                tables.append(" | ".join(cell.text for cell in row.cells))

        return "\n".join(paragraphs + tables)
    except Exception as exc:
        raise ExtractionError(f"Gagal membaca file DOCX: {exc}") from exc


def _ocr_image(path: Path) -> str:
    try:
        with Image.open(path) as image:
            return _safe_ocr(image)
    except Exception as exc:
        raise ExtractionError(f"Gagal membaca file gambar: {exc}") from exc


